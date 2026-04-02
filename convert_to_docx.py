#!/usr/bin/env python3
"""Конвертация .md файлов в .docx с нативными формулами Word (OMML)."""

import re
import os
from lxml import etree
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import latex2mathml.converter

BASE = r"C:\Users\Ruslan\Проекты\тпр"
XSLT_PATH = r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL"

FILES = [
    ("лекции_ТПР.md", "лекции_ТПР.docx"),
    ("лекции_ТПР_оригинал.md", "лекции_ТПР_оригинал.docx"),
    ("ответы_экзамен_ТПР.md", "ответы_экзамен_ТПР.docx"),
]

# Load XSLT transform once
xslt_tree = etree.parse(XSLT_PATH)
xslt_transform = etree.XSLT(xslt_tree)


def latex_to_omml(latex_str):
    """Convert LaTeX formula to Word OMML XML element."""
    try:
        # LaTeX -> MathML
        mathml_str = latex2mathml.converter.convert(latex_str)
        # MathML -> OMML via XSLT
        mathml_tree = etree.fromstring(mathml_str.encode('utf-8'))
        omml_tree = xslt_transform(mathml_tree)
        return omml_tree.getroot()
    except Exception as e:
        # Fallback: return latex as plain text
        return None


def set_cell_shading(cell, color):
    """Set cell background color."""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    shading.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading)


def add_formatted_text(paragraph, text):
    """Add text with inline formatting (bold, italic, code, LaTeX formulas)."""
    # Split by: $$...$$ (display), $...$ (inline), **...**, *...*, `...`
    pattern = r'(\$\$[^$]+\$\$|\$[^$]+\$|\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)'
    parts = re.split(pattern, text)
    for part in parts:
        if not part:
            continue
        # Display math $$...$$
        if part.startswith('$$') and part.endswith('$$'):
            latex = part[2:-2].strip()
            omml = latex_to_omml(latex)
            if omml is not None:
                paragraph._p.append(omml)
            else:
                run = paragraph.add_run(f' {latex} ')
                run.font.name = 'Cambria Math'
                run.italic = True
            continue
        # Inline math $...$
        if part.startswith('$') and part.endswith('$') and len(part) > 2:
            latex = part[1:-1].strip()
            omml = latex_to_omml(latex)
            if omml is not None:
                paragraph._p.append(omml)
            else:
                run = paragraph.add_run(f' {latex} ')
                run.font.name = 'Cambria Math'
                run.italic = True
            continue
        if part.startswith('***') and part.endswith('***'):
            run = paragraph.add_run(part[3:-3])
            run.bold = True
            run.italic = True
        elif part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x80, 0x00, 0x00)
        else:
            paragraph.add_run(part)


def parse_table(lines):
    """Parse markdown table lines into rows of cells."""
    rows = []
    for line in lines:
        line = line.strip()
        if line.startswith('|'):
            line = line[1:]
        if line.endswith('|'):
            line = line[:-1]
        cells = [c.strip() for c in line.split('|')]
        if all(re.match(r'^[-:]+$', c) for c in cells if c):
            continue
        rows.append(cells)
    return rows


def is_display_math_line(stripped):
    """Check if line is a standalone display math block $$...$$."""
    return stripped.startswith('$$') and stripped.endswith('$$') and len(stripped) > 4


def md_to_docx(md_path, docx_path):
    """Convert a markdown file to .docx with formatting and native Word formulas."""
    print(f"  Конвертация: {os.path.basename(md_path)} -> {os.path.basename(docx_path)}")

    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    # Set heading styles
    for i in range(1, 5):
        hs = doc.styles[f'Heading {i}']
        hs.font.name = 'Times New Roman'
        hs.font.color.rgb = RGBColor(0, 0, 0)
        if i == 1:
            hs.font.size = Pt(18)
        elif i == 2:
            hs.font.size = Pt(15)
        elif i == 3:
            hs.font.size = Pt(13)
        else:
            hs.font.size = Pt(12)

    lines = content.split('\n')
    i = 0
    formula_ok = 0
    formula_fallback = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Skip empty lines
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if re.match(r'^-{3,}$', stripped) or re.match(r'^\*{3,}$', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), 'auto')
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        # Headings
        heading_match = re.match(r'^(#{1,4})\s+(.+)$', stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            heading = doc.add_heading(level=level)
            add_formatted_text(heading, text)
            i += 1
            continue

        # Display math on its own line: $$...$$
        if is_display_math_line(stripped):
            latex = stripped[2:-2].strip()
            p = doc.add_paragraph()
            p.alignment = 1  # CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            omml = latex_to_omml(latex)
            if omml is not None:
                p._p.append(omml)
                formula_ok += 1
            else:
                run = p.add_run(latex)
                run.font.name = 'Cambria Math'
                run.italic = True
                run.font.size = Pt(12)
                formula_fallback += 1
            i += 1
            continue

        # Multi-line display math: $$ on one line, content, $$ on another
        if stripped == '$$':
            math_lines = []
            i += 1
            while i < len(lines) and lines[i].strip() != '$$':
                math_lines.append(lines[i])
                i += 1
            i += 1  # skip closing $$
            latex = '\n'.join(math_lines).strip()
            p = doc.add_paragraph()
            p.alignment = 1  # CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            omml = latex_to_omml(latex)
            if omml is not None:
                p._p.append(omml)
                formula_ok += 1
            else:
                run = p.add_run(latex)
                run.font.name = 'Cambria Math'
                run.italic = True
                formula_fallback += 1
            continue

        # Code blocks
        if stripped.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1

            code_text = '\n'.join(code_lines)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.left_indent = Cm(1)
            run = p.add_run(code_text)
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), 'F0F0F0')
            shading.set(qn('w:val'), 'clear')
            run._r.get_or_add_rPr().append(shading)
            continue

        # Tables
        if stripped.startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i])
                i += 1

            rows = parse_table(table_lines)
            if not rows:
                continue

            num_cols = max(len(r) for r in rows)
            for r in rows:
                while len(r) < num_cols:
                    r.append('')

            table = doc.add_table(rows=len(rows), cols=num_cols)
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            for ri, row in enumerate(rows):
                for ci, cell_text in enumerate(row):
                    cell = table.rows[ri].cells[ci]
                    cell.text = ''
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    add_formatted_text(p, cell_text)
                    if ri == 0:
                        set_cell_shading(cell, 'D9E2F3')
                        for run in p.runs:
                            run.bold = True
            continue

        # Ordered list
        ol_match = re.match(r'^(\d+)\.\s+(.+)$', stripped)
        if ol_match:
            p = doc.add_paragraph(style='List Number')
            add_formatted_text(p, ol_match.group(2))
            i += 1
            continue

        # Unordered list
        ul_match = re.match(r'^[-*]\s+(.+)$', stripped)
        if ul_match:
            p = doc.add_paragraph(style='List Bullet')
            add_formatted_text(p, ul_match.group(1))
            i += 1
            continue

        # Sub-list items (indented)
        sub_match = re.match(r'^  [-*]\s+(.+)$', line)
        if sub_match:
            p = doc.add_paragraph(style='List Bullet 2')
            add_formatted_text(p, sub_match.group(1))
            i += 1
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        add_formatted_text(p, stripped)
        i += 1

    doc.save(docx_path)
    print(f"    Сохранено: {docx_path}")
    print(f"    Формулы OMML: {formula_ok}, fallback (текст): {formula_fallback}")


def main():
    print("=" * 60)
    print("Конвертация MD -> DOCX (с нативными формулами Word)")
    print("=" * 60)

    for md_name, docx_name in FILES:
        md_path = os.path.join(BASE, md_name)
        docx_path = os.path.join(BASE, docx_name)

        if not os.path.exists(md_path):
            print(f"  ПРОПУСК: {md_name} не найден")
            continue

        md_to_docx(md_path, docx_path)

    print("\nГотово!")


if __name__ == '__main__':
    main()
